"""Document loaders: one loader per filetype, returning normalized `Document` objects.

Each loader produces one `Document` per logical segment (e.g. a PDF page). Downstream
chunking subdivides those Documents further while preserving provenance metadata.

Adding a new loader:
1. Subclass `BaseLoader`, set `extensions`, implement `load`.
2. Register it in `LOADER_REGISTRY` (or via `register_loader`).
"""

from __future__ import annotations

from abc import ABC, abstractmethod
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, ClassVar

from docx import Document as DocxDocument
from pypdf import PdfReader


@dataclass(frozen=True)
class Document:
    """A loaded document segment with provenance metadata.

    `content` is raw extracted text. `metadata` carries at minimum `source_file`
    and (when applicable) `page`. Chunking will add `chunk_index` and `chunk_id`.
    """

    content: str
    metadata: dict[str, Any] = field(default_factory=dict)


class BaseLoader(ABC):
    """Abstract base for filetype-specific loaders."""

    extensions: ClassVar[tuple[str, ...]] = ()

    @abstractmethod
    def load(self, path: Path) -> list[Document]:
        """Load `path` and return one `Document` per logical segment."""


class PDFLoader(BaseLoader):
    """Load PDF files page by page using pypdf."""

    extensions: ClassVar[tuple[str, ...]] = (".pdf",)

    def load(self, path: Path) -> list[Document]:
        """Extract text per page; empty pages produce empty Documents so page numbers align."""
        reader = PdfReader(str(path))
        source = path.as_posix()
        docs: list[Document] = []
        for i, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            docs.append(
                Document(
                    content=text.strip(),
                    metadata={"source_file": source, "page": i},
                )
            )
        return docs


class TextLoader(BaseLoader):
    """Load plain text files as a single Document."""

    extensions: ClassVar[tuple[str, ...]] = (".txt",)

    def load(self, path: Path) -> list[Document]:
        """Read the file as UTF-8 text and return one Document."""
        return [
            Document(
                content=path.read_text(encoding="utf-8").strip(),
                metadata={"source_file": path.as_posix()},
            )
        ]


class MarkdownLoader(BaseLoader):
    """Load Markdown files as a single Document, preserving raw markdown.

    The LLM is fine reading markdown directly; the chunker splits on paragraph
    breaks (double newlines) which works well for markdown sections.
    """

    extensions: ClassVar[tuple[str, ...]] = (".md", ".markdown")

    def load(self, path: Path) -> list[Document]:
        """Read the markdown file as UTF-8 text and return one Document."""
        return [
            Document(
                content=path.read_text(encoding="utf-8").strip(),
                metadata={"source_file": path.as_posix()},
            )
        ]


class DOCXLoader(BaseLoader):
    """Load .docx files via python-docx, concatenating non-empty paragraphs."""

    extensions: ClassVar[tuple[str, ...]] = (".docx",)

    def load(self, path: Path) -> list[Document]:
        """Concatenate paragraph text with double newlines between paragraphs."""
        doc = DocxDocument(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
        content = "\n\n".join(paragraphs)
        return [
            Document(
                content=content,
                metadata={"source_file": path.as_posix()},
            )
        ]


LOADER_REGISTRY: dict[str, type[BaseLoader]] = {
    ".pdf": PDFLoader,
    ".txt": TextLoader,
    ".md": MarkdownLoader,
    ".markdown": MarkdownLoader,
    ".docx": DOCXLoader,
}


def register_loader(loader_cls: type[BaseLoader]) -> None:
    """Register `loader_cls` against each of its declared extensions."""
    for ext in loader_cls.extensions:
        LOADER_REGISTRY[ext.lower()] = loader_cls


def get_loader(path: Path) -> BaseLoader:
    """Return a loader instance for `path` based on its file extension.

    Raises `ValueError` if no loader is registered for the extension.
    """
    ext = path.suffix.lower()
    loader_cls = LOADER_REGISTRY.get(ext)
    if loader_cls is None:
        supported = sorted(LOADER_REGISTRY.keys())
        raise ValueError(f"No loader registered for extension {ext!r}. Supported: {supported}")
    return loader_cls()
